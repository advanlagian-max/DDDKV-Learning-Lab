#!/usr/bin/env python3
# gen-matriks.py — Membuat matriks 36 minggu (72 pertemuan) dari data kurikulum terverifikasi.
# Keluaran: docs/06-matriks-36-minggu.md , .csv , .xlsx
# Pemakaian: python3 tools/gen-matriks.py   (jalankan dari root folder dkv-learning-lab)
import json, os, csv, re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "app", "data")
DOCS = os.path.join(ROOT, "docs")

BUKU = {
    "B1": "Buku Siswa Sem 1 (Desain Komunikasi Visual Kls X Sem 1)",
    "B2": "Buku Siswa Sem 2 (Dasar-Dasar DKV Kls X Sem 2)",
    "SUP": "Dokumen Instrumen Supervisi Akademik (SUP)",
}

# Simulasi: mana yang benar-benar berfungsi di aplikasi (hasil pengujian 13 Okt 2026)
SIM_FUNGSI = {"sim-tipografi", "sim-fotografi", "sim-bitmap-vektor", "sim-alur-produksi"}

STATUS_SUMBER = {"sesuai-buku": "sesuai buku", "sebagian": "sebagian buku + usulan"}

def ringkas(s, n=200):
    s = re.sub(r"\s+", " ", str(s)).strip()
    return s if len(s) <= n else s[: n - 1] + "…"

def muat():
    minggu = []
    for f, sem in [("kurikulum-sem1.json", 1), ("kurikulum-sem2.json", 2)]:
        d = json.load(open(os.path.join(DATA, f), encoding="utf-8"))
        for w in d["minggu"]:
            w["_sem"] = sem
            minggu.append(w)
    return minggu

def sel_sumber(w):
    baris = []
    for s in w["sumber"]:
        kode = s.get("kode", "")
        nama_kode = {"SUP": "Instrumen Supervisi Akademik (aspek C.1–C.4)"}.get(kode, kode)
        bagian = " · ".join(x for x in [s.get("bab") or "", s.get("subbab") or ""] if x)
        hal = f"hal. {s['halaman']}" if s.get("halaman") else ""
        isi = f"{nama_kode} {bagian}".strip()
        baris.append(f"〔B〕 {isi} ({hal})" if hal else f"〔B〕 {isi}")
    st = STATUS_SUMBER.get(w.get("statusSumber"), w.get("statusSumber") or "-")
    baris.append(f"Status sumber: **{st}**")
    if w.get("catatanVerifikasi"):
        baris.append(f"⚠ {ringkas(w['catatanVerifikasi'], 160)}")
    return "<br>".join(baris)

def sel_tujuan(w):
    isi = " ".join(f"{i+1}) {ringkas(t,190)}" for i, t in enumerate(w["tujuan"]))
    return f"⧗ *usulan (CP/ATP belum diverifikasi sekolah)*: {isi}"

def sel_konsep(w):
    return "<br>".join(f"〔B〕 {ringkas(x, 210)}" for x in w["materiInti"])

def sel_akt(w, kunci):
    return "; ".join(f"{i+1}) {ringkas(x,130)}" for i, x in enumerate(w["aktivitas"].get(kunci, [])))

def sel_simulasi(w):
    s = w.get("simulasi") or {}
    if not s:
        return "tidak dijadwalkan simulasi digital → praktik langsung di kelas 〔U〕"
    tag = "〔F〕 berfungsi di aplikasi" if s.get("id") in SIM_FUNGSI else "〔T7〕 masih dalam pengembangan (ditandai di aplikasi)"
    return f"{s.get('judul','')} (`{s.get('id','')}`) — {tag}"

def sel_asesmen(w):
    a = w.get("asesmenFormatif") or {}
    bentuk = str(a.get("bentuk", ""))
    label = {"kuis-pilihan-ganda": "Kuis pilihan ganda", "asesmen-sumatif-bab-dan-semester": "Asesmen sumatif BAB & semester",
             "uji-kompetensi-ebab-buku": "Uji Kompetensi buku (BAB 1 Tipografi)",
             "asesmen-sumatif-akhir-semester": "Asesmen sumatif akhir semester"}.get(bentuk, bentuk or "asesmen formatif")
    isi = f"Formatif: {label} — {a.get('jumlahSoal','?')} soal; cakupan: {ringkas(a.get('cakupan',''), 170)} 〔U〕"
    if "sumatif" in bentuk or "uji-kompetensi" in bentuk:
        isi += "<br>**＋ pekan asesmen sumatif/uji kompetensi**"
    return isi

def sel_portofolio(w):
    if w.get("proyek"):
        p = w["proyek"]
        return f"〔U〕 Proyek: {p.get('judul','')} ({p.get('jenis','')}) — luaran: {ringkas(p.get('luaran',''),150)}; tenggat {p.get('tenggat','')} + refleksi minggu ini"
    lk = w.get("lkpd") or {}
    return f"〔U〕 LKPD: {lk.get('judul','-')} (lembar kerja terisi) + refleksi minggu ini"

def topik_dua(w):
    return f"**Senin:** {w['pertemuanSenin']['tema']}<br>**Rabu:** {w['pertemuanRabu']['tema']}"

KEPALA = ["Minggu", "Topik (Senin / Rabu)", "Sumber & bab terkait", "Tujuan pembelajaran (usulan)", "Konsep penting (dari buku)",
          "Aktivitas memahami", "Aktivitas mengaplikasi", "Aktivitas merefleksi", "Simulasi/praktik", "Asesmen", "Bukti portofolio"]

def baris_matrik(w):
    return [f"**{w['minggu']}** (Sm {w['_sem']})", topik_dua(w), sel_sumber(w), sel_tujuan(w), sel_konsep(w),
            sel_akt(w, "memahami"), sel_akt(w, "mengaplikasi"), sel_akt(w, "merefleksi"),
            sel_simulasi(w), sel_asesmen(w), sel_portofolio(w)]

def tabel(minggu):
    out = ["| " + " | ".join(KEPALA) + " |", "|" + "---|" * len(KEPALA)]
    for w in minggu:
        out.append("| " + " | ".join(baris_matrik(w)) + " |")
    return "\n".join(out)

def detail_pertemuan(minggu):
    out = []
    for w in minggu:
        for hari, kunci in [("Senin", "pertemuanSenin"), ("Rabu", "pertemuanRabu")]:
            p = w[kunci]
            total = sum(a.get("menit", 0) for a in p["agenda"])
            out.append(f"### Minggu {w['minggu']} · {hari} — {p['tema']}")
            out.append(f"{p['jp']} JP · {p['menit']} menit (rincian agenda {total} menit termasuk transisi/istirahat · 1 JP = 45 menit) 〔U〕\n")
            out.append("| Tahap | Menit | Kegiatan |")
            out.append("|---|---|---|")
            for a in p["agenda"]:
                out.append(f"| {a.get('tahap','')} | {a.get('menit','')} | {ringkas(a.get('kegiatan',''), 320)} |")
            out.append("")
    return "\n".join(out)

def rekap(minggu):
    jp = sum(w["pertemuanSenin"]["jp"] + w["pertemuanRabu"]["jp"] for w in minggu)
    menit = sum(w["pertemuanSenin"]["menit"] + w["pertemuanRabu"]["menit"] for w in minggu)
    return len(minggu), len(minggu) * 2, jp, menit

def tulis_md(minggu):
    n_w, n_p, jp, menit = rekap(minggu)
    sem1 = [w for w in minggu if w["_sem"] == 1]
    sem2 = [w for w in minggu if w["_sem"] == 2]
    verif = [w for w in minggu if w.get("catatanVerifikasi")]
    L = []
    L.append("# 06 — Matriks 36 Minggu & 72 Pertemuan (Analisis Buku DD DKV + Instrumen Supervisi)\n")
    L.append("Dasar: **Buku Siswa Dasar-Dasar Desain Komunikasi Visual Kelas X Semester 1 & 2 (Kemendikbudristek, 2022)** dan **Dokumen Instrumen Supervisi Akademik** sekolah. "
             "Disusun otomatis dari data kurikulum terverifikasi (`app/data/kurikulum-sem1.json`, `kurikulum-sem2.json`) memakai `tools/gen-matriks.py`.\n")
    L.append(f"Ringkasan: **{n_w} minggu · {n_p} pertemuan · {jp} JP · {menit:,} menit** "
             f"(setiap pertemuan 6 JP = 270 menit; 1 JP = 45 menit; 270 menit sudah termasuk transisi/istirahat).\n".replace(",", "."))
    L.append("> **Catatan kejujuran dokumen:** matriks ini adalah **rancangan operasional pembelajaran**, bukan ATP resmi dan bukan pengganti isi buku. "
             "Bagian yang bersumber dari buku ditandai `〔B〕` beserta bab/halaman; usulan desain pembelajaran ditandai `〔U〕`; "
             "hal yang harus diverifikasi terhadap dokumen resmi sekolah ditandai `⧗`; simulasi yang belum selesai ditandai `〔T7〕`.\n")
    L.append("---\n")

    L.append("## 1. Cara membaca tanda\n")
    L.append("| Tanda | Arti |\n|---|---|")
    L.append("| 〔B〕 | **Bersumber buku** — teks/konsep/kegiatan berasal dari buku siswa; bab & halaman dicantumkan |")
    L.append("| 〔U〕 | **Usulan desain pembelajaran** — rumusan guru/penyusun, tidak diklaim sebagai isi buku |")
    L.append("| ⧗ | **Perlu diverifikasi** dengan dokumen resmi sekolah (CP, ATP, indikator, kalender) |")
    L.append("| ⚠ | **Isi/kata kunci perlu dikonfirmasi guru** sebelum dipakai (mis. istilah, lisensi perangkat lunak) |")
    L.append("| 〔F〕 | Simulasi digital **sudah berfungsi** dan sudah diuji di aplikasi |")
    L.append("| 〔T7〕 | Simulasi digital **belum selesai** (ditandai “dalam pengembangan” pada aplikasi) |")
    L.append("")
    L.append("Kolom **Sumber & bab terkait** dan **Konsep penting** = isi buku (〔B〕). "
             "Kolom **Tujuan pembelajaran, Aktivitas memahami/mengaplikasi/merefleksi, Asesmen, dan Bukti portofolio** = usulan desain pembelajaran (〔U〕) "
             "dengan tetap berpijak pada aktivitas buku (mis. “Ayo Berdiskusi”, “Ayo Praktikkan”, “Ayo Berkreasi”).\n")

    L.append("## 2. Ringkasan analisis sumber\n")
    L.append("### 2.1 Buku Semester 1 (kode B1) — bab dan halaman cetak\n")
    L.append("| Bab/bagian | Halaman cetak | Dipakai pada |\n|---|---|---|")
    for b, h, m in [("Petunjuk Penggunaan Buku", "xii", "—"),
                    ("BAB 1 Profil Technopreneur, Peluang Usaha, dan Profesi di Bidang DKV", "1–16", "Minggu 1–4"),
                    ("BAB 2 Keutamaan Keselamatan dan Kesehatan Kerja", "17–50", "Minggu 5–7"),
                    ("BAB 3 Perkembangan Teknologi Industri dan Dunia Kerja Bidang DKV", "51–74", "Minggu 8–11"),
                    ("BAB 4 Teknik Dasar Proses Produksi (Kreativitas) pada Industri DKV", "75–92", "Minggu 12–15"),
                    ("BAB 5 Sketsa dan Ilustrasi", "93–132", "Minggu 16–18"),
                    ("Latihan Soal Akhir Semester 1", "133–144", "Minggu 18"),
                    ("Glosarium", "≥145", "rujukan istilah")]:
        L.append(f"| {b} | {h} | {m} |")
    L.append("\n### 2.2 Buku Semester 2 (kode B2) — bab dan halaman cetak\n")
    L.append("| Bab/bagian | Halaman cetak | Dipakai pada |\n|---|---|---|")
    for b, h, m in [("Bab 1 Tipografi (sejarah, jenis, anatomi, prinsip, leading/kerning/tracking)", "1–30", "Minggu 19–24"),
                    ("Bab 2 Fotografi Dasar (kamera, segitiga eksposur, ukuran gambar, sudut, pencahayaan, penyuntingan)", "31–86", "Minggu 25–30"),
                    ("Bab 3 Komputer Grafis (bitmap/vektor, perangkat lunak, praktik, produk)", "87–142", "Minggu 31–36"),
                    ("Latihan Soal Akhir Semester 2", "143–156", "Minggu 36"),
                    ("Glosarium", "≥157", "rujukan istilah")]:
        L.append(f"| {b} | {h} | {m} |")
    L.append("\n**Empat produk pada Bab 3.C** (dirujuk minggu 35): mug (8 × 20 cm, dibuat dengan CorelDRAW X7 pada buku), logo, packaging, kartu nama. "
             "〔B〕 h. 135–137. Karena buku menyebut perangkat lunak berlisensi, praktik sekolah diarahkan ke alternatif bebas (Inkscape/GIMP) ⚠.\n")
    L.append("### 2.3 Instrumen Supervisi Akademik (kode SUP)\n")
    L.append("Skala **1–4**; aspek yang dinilai beserta sub-butir yang dipakai aplikasi:\n")
    L.append("| Aspek | Sub-butir pada matriks |\n|---|---|")
    L.append("| C.1 Perencanaan | tujuan jelas & terukur; materi sesuai CP/materi buku; strategi pembelajaran mendalam; diferensiasi |")
    L.append("| C.2 Pelaksanaan | pendahuluan/apersepsi; interaksi guru–murid; aktivitas murid; kolaborasi; pemanfaatan teknologi & sumber belajar |")
    L.append("| C.3 Penerapan Deep Learning (5C) | Critical Thinking, Creativity, Collaboration, Communication, Connection to Real Life |")
    L.append("| C.4 Asesmen & Refleksi | asesmen formatif berjalan; sumatif/praktik; refleksi guru; refleksi murid |")
    L.append("\nCatatan instrumen: integrasi program sekolah (Adiwiyata, Gapura Pancawaluya, Gerakan Indonesia ASRI, Literasi & Numerasi, S3, SimPel) ke dalam "
             "modul ajar dan pelaksanaan pembelajaran. Kolom `catatanSupervisi` pada tiap minggu memetakan rencana ke aspek C.1–C.4 (kata kunci: perencanaan, "
             "pelaksanaan, asesmen, deepLearning/mindful). Tanda tangan dokumen asli: Saidun, S.Pd., M.M. (NIP 197009201994021002).\n")
    L.append("### 2.4 Status CP, ATP, dan kalender ⧗\n")
    L.append("- Buku siswa **tidak memuat rumusan CP/ATP**; seluruh tujuan pembelajaran di matriks ini adalah **usulan** yang menunggu verifikasi terhadap dokumen kurikulum resmi sekolah ⧗.\n"
             "- Basis kurikulum yang dirujuk: Kepmen BSKAP No. 008/H/KR/2022 dan 033/H/KR/2022 (rujukan dokumen resmi) — **kutipan pasalnya tidak diklaim** di aplikasi ini.\n"
             "- Kalender pendidikan sekolah belum tersedia; minggu efektif & pekan asesmen pada matriks masih asumsi (Sem 1 minggu 1–18, Sem 2 minggu 19–36) ⧗.\n"
             "- Validasi silang rujukan halaman: seluruh judul subbab pada kolom sumber dicek otomatis terhadap teks hasil ekstraksi buku "
             "(halaman cetak → berkas PDF: Sem 1 +12, Sem 2 +14). Hasil: **semua rujukan halaman cocok**; selisih awal hanya karena penyingkatan judul subbab.\n")

    L.append("---\n")
    L.append("## 3. Matriks Semester 1 (minggu 1–18)\n")
    L.append("Setiap baris = 1 minggu = 2 pertemuan (Senin & Rabu), masing-masing 6 JP / 270 menit.\n")
    L.append(tabel(sem1))
    L.append("\n## 4. Matriks Semester 2 (minggu 19–36)\n")
    L.append("Setiap baris = 1 minggu = 2 pertemuan (Senin & Rabu), masing-masing 6 JP / 270 menit.\n")
    L.append(tabel(sem2))

    L.append("\n---\n")
    L.append("## 5. Rincian 72 pertemuan (Senin 6 JP/270 menit · Rabu 6 JP/270 menit)\n")
    L.append("Agenda di bawah adalah rancangan 〔U〕 yang sudah tertanam pada data kurikulum; setiap pertemuan berjumlah 6 JP = 270 menit (termasuk transisi/istirahat).\n")
    L.append(detail_pertemuan(minggu))

    L.append("---\n## 6. Peta simulasi (nyata vs dalam pengembangan)\n")
    L.append("| Simulasi | Minggu pemakaian | Status jujur |\n|---|---|---|")
    for sid in ["sim-technopreneur", "sim-k3-ergonomi", "sim-bitmap-vektor", "sim-alur-produksi", "sim-sketsa-lab",
                "sim-tipografi", "sim-fotografi", "sim-studio-dkv"]:
        ms = [w["minggu"] for w in minggu if (w.get("simulasi") or {}).get("id") == sid]
        judul = next((w["simulasi"]["judul"] for w in minggu if (w.get("simulasi") or {}).get("id") == sid), sid)
        status = "**berfungsi & sudah diuji** 〔F〕" if sid in SIM_FUNGSI else "**dalam pengembangan** 〔T7〕 — ditandai di antarmuka, bukan tombol palsu"
        L.append(f"| {judul} | {', '.join(str(x) for x in ms) or '—'} | {status} |")
    L.append("\nMinggu tanpa simulasi digital: " + ", ".join(str(w["minggu"]) for w in minggu if not w.get("simulasi")) +
             " → diisi praktik langsung (latihan gambar/observasi/diskusi kelas) 〔U〕.\n")
    L.append("Semua simulasi berlabel: **“simulasi edukatif, bukan pengganti perangkat lunak profesional”**.\n")

    L.append("---\n## 7. Peta asesmen & bukti portofolio\n")
    L.append("| Minggu | Asesmen formatif (usulan) | Pekan sumatif | Bukti portofolio |\n|---|---|---|---|")
    for w in minggu:
        a = w.get("asesmenFormatif") or {}
        b = str(a.get("bentuk", ""))
        summ = "ya ⧗" if ("sumatif" in b or "uji-kompetensi" in b) else "—"
        L.append(f"| {w['minggu']} | {b} ({a.get('jumlahSoal','?')} soal) — {ringkas(a.get('cakupan',''),120)} | {summ} | {sel_portofolio(w)} |")
    L.append("\nAsesmen diagnostik: minggu 1 (kuis awal). Asesmen sumatif/uji kompetensi: " +
             ", ".join(f"minggu {w['minggu']}" for w in minggu if "sumatif" in str((w.get('asesmenFormatif') or {}).get('bentuk','')) or "uji-kompetensi" in str((w.get('asesmenFormatif') or {}).get('bentuk',''))) +
             ". Penilaian karya visual/terbuka dilakukan guru (rubrik per minggu); kuis objektif dinilai otomatis di server.\n")

    L.append("---\n## 8. Rekap waktu\n")
    L.append("| Item | Nilai |\n|---|---|")
    L.append(f"| Minggu | {n_w} (18 minggu per semester) |")
    L.append(f"| Pertemuan | {n_p} (2 per minggu: Senin & Rabu) |")
    L.append(f"| JP per pertemuan | 6 JP (270 menit; 1 JP = 45 menit) |")
    L.append(f"| Total JP | {jp} JP |")
    L.append(f"| Total menit | {menit} menit = {menit/60:.0f} jam pelaksanaan |")
    L.append("\nRincian per minggu sama untuk seluruh semester (6 JP Senin + 6 JP Rabu); guru dapat mengubah hari, jumlah JP, durasi JP, dan minggu efektif "
             "melalui menu pengaturan aplikasi tanpa menyunting kode.\n")

    L.append("---\n## 9. Daftar yang wajib diverifikasi ⧗\n")
    L.append("1. **CP/ATP** seluruh 36 minggu — buku tidak memuatnya; bandingkan dengan CP Fase E & ATP resmi sekolah.\n"
             "2. **Kalender pendidikan** — minggu efektif, pekan asesmen, dan hari libur.\n"
             "3. **Indikator & bobot supervisi** — samakan dengan Instrumen Supervisi yang berlaku (skor 1–4, aspek C.1–C.4).\n"
             "4. **Lisensi perangkat lunak** — CorelDRAW/Adobe pada buku; alternatif Inkscape/GIMP ⚠.\n"
             "5. Butir-butir berikut ditandai `⚠` pada matriks karena isinya paling perlu dicek guru:\n")
    for w in verif:
        L.append(f"   - **Minggu {w['minggu']}** ({w['tema']}): {ringkas(w['catatanVerifikasi'], 260)}")
    L.append("")
    open(os.path.join(DOCS, "06-matriks-36-minggu.md"), "w", encoding="utf-8").write("\n".join(L))
    return n_w, n_p, jp, menit

def tulis_csv(minggu):
    path = os.path.join(DOCS, "06-matriks-36-minggu.csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(["semester", "minggu", "topik_senin", "topik_rabu", "senin_jp", "senin_menit", "rabu_jp", "rabu_menit",
                     "sumber_dan_bab", "status_sumber", "tujuan_usulan", "konsep_penting_buku", "memahami", "mengaplikasi",
                     "merefleksi", "simulasi", "asesmen", "bukti_portofolio", "catatan_verifikasi"])
        for w in minggu:
            wr.writerow([w["_sem"], w["minggu"], w["pertemuanSenin"]["tema"], w["pertemuanRabu"]["tema"],
                         w["pertemuanSenin"]["jp"], w["pertemuanSenin"]["menit"], w["pertemuanRabu"]["jp"], w["pertemuanRabu"]["menit"],
                         " | ".join(f"{s.get('kode')} {s.get('bab','')} {s.get('subbab','')} (hal. {s.get('halaman','-')})" for s in w["sumber"]),
                         STATUS_SUMBER.get(w.get("statusSumber"), "-"),
                         " ".join(f"{i+1}) {t}" for i, t in enumerate(w["tujuan"])),
                         " || ".join(w["materiInti"]),
                         " || ".join(w["aktivitas"].get("memahami", [])),
                         " || ".join(w["aktivitas"].get("mengaplikasi", [])),
                         " || ".join(w["aktivitas"].get("merefleksi", [])),
                         re.sub(r"<br>", " | ", sel_simulasi(w)),
                         re.sub(r"<br>", " | ", sel_asesmen(w)),
                         re.sub(r"<br>", " | ", sel_portofolio(w)),
                         w.get("catatanVerifikasi") or ""])
    return path

def tulis_xlsx(minggu):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except Exception:
        return None
    wb = Workbook()
    kepala_fill = PatternFill("solid", fgColor="1E3A8A")
    kepala_font = Font(color="FFFFFF", bold=True)
    sempit = Alignment(wrap_text=True, vertical="top")

    def lembar(nama, kepala, baris):
        ws = wb.create_sheet(nama)
        ws.append(kepala)
        for c in ws[1]:
            c.fill, c.font, c.alignment = kepala_fill, kepala_font, sempit
        for r in baris:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = sempit
        for i, k in enumerate(kepala, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 34 if "Konsep" in k or "Sumber" in k else 26
        ws.freeze_panes = "A2"
        return ws

    for sem, nama in [(1, "Matriks Sem 1"), (2, "Matriks Sem 2")]:
        lembar(nama, KEPALA, [baris_matrik(w) for w in minggu if w["_sem"] == sem])

    baris_pertemuan = []
    for w in minggu:
        for hari, kunci in [("Senin", "pertemuanSenin"), ("Rabu", "pertemuanRabu")]:
            p = w[kunci]
            baris_pertemuan.append([w["_sem"], w["minggu"], hari, p["tema"], p["jp"], p["menit"],
                                    " → ".join(f"{a.get('tahap','')} ({a.get('menit','')}′)" for a in p["agenda"]),
                                    " | ".join(a.get("kegiatan", "") for a in p["agenda"])])
    lembar("Rincian 72 Pertemuan", ["Semester", "Minggu", "Hari", "Tema pertemuan", "JP", "Menit", "Alur tahap (menit)", "Rincian kegiatan"], baris_pertemuan)

    sim = []
    for sid in ["sim-technopreneur", "sim-k3-ergonomi", "sim-bitmap-vektor", "sim-alur-produksi", "sim-sketsa-lab", "sim-tipografi", "sim-fotografi", "sim-studio-dkv"]:
        ms = [w["minggu"] for w in minggu if (w.get("simulasi") or {}).get("id") == sid]
        judul = next((w["simulasi"]["judul"] for w in minggu if (w.get("simulasi") or {}).get("id") == sid), sid)
        sim.append([sid, judul, ", ".join(map(str, ms)) or "—", "berfungsi (sudah diuji)" if sid in SIM_FUNGSI else "dalam pengembangan 〔T7〕"])
    lembar("Simulasi", ["ID", "Judul simulasi", "Minggu", "Status jujur"], sim)

    ase = [[w["_sem"], w["minggu"], (w.get("asesmenFormatif") or {}).get("bentuk", ""), (w.get("asesmenFormatif") or {}).get("jumlahSoal", ""),
            (w.get("asesmenFormatif") or {}).get("cakupan", ""), re.sub(r"<br>", " | ", sel_portofolio(w))]
           for w in minggu]
    lembar("Asesmen-Portofolio", ["Semester", "Minggu", "Bentuk asesmen", "Jumlah soal", "Cakupan", "Bukti portofolio (usulan)"], ase)

    vf = [[w["minggu"], w["tema"], w.get("catatanVerifikasi") or ""] for w in minggu if w.get("catatanVerifikasi")]
    vf.append(["—", "CP & ATP seluruh minggu", "Buku tidak memuat CP/ATP → usulan perlu diverifikasi terhadap dokumen resmi sekolah ⧗"])
    vf.append(["—", "Kalender pendidikan", "Minggu efektif & pekan asesmen masih asumsi; kalender resmi belum tersedia ⧗"])
    vf.append(["—", "Indikator supervisi", "Samakan istilah/bobot dengan Instrumen Supervisi yang berlaku (C.1–C.4, skor 1–4) ⧗"])
    lembar("Catatan Verifikasi", ["Minggu", "Tema/catatan", "Butir yang perlu diverifikasi ⧗/⚠"], vf)

    del wb["Sheet"]
    ws0 = wb.create_sheet("Ringkasan", 0)
    n_w, n_p, jp, menit = rekap(minggu)
    for r in [["DKV Learning Lab — Matriks 36 Minggu / 72 Pertemuan", ""],
              ["Minggu", n_w], ["Pertemuan", n_p], ["JP per pertemuan", "6 JP (270 menit)"],
              ["Total JP", jp], ["Total menit", menit], ["Sumber data", "app/data/kurikulum-sem1.json + kurikulum-sem2.json"],
              ["Buku acuan", "B1 Buku Sem 1 (176 hal) · B2 Buku Sem 2 (184 hal) · SUP Instrumen Supervisi"],
              ["Tanda", "〔B〕 dari buku · 〔U〕 usulan desain · ⧗ perlu verifikasi · ⚠ perlu konfirmasi guru · 〔F〕 simulasi berfungsi · 〔T7〕 simulasi dalam pengembangan"]]:
        ws0.append(r)
    ws0.column_dimensions["A"].width = 32
    ws0.column_dimensions["B"].width = 90
    ws0["A1"].font = Font(bold=True, size=13)
    path = os.path.join(DOCS, "06-matriks-36-minggu.xlsx")
    wb.save(path)
    return path

if __name__ == "__main__":
    mg = muat()
    info = tulis_md(mg)
    p1 = tulis_csv(mg)
    p2 = tulis_xlsx(mg)
    print(f"MD  : docs/06-matriks-36-minggu.md  ({info[0]} minggu, {info[1]} pertemuan, {info[2]} JP, {info[3]} menit)")
    print(f"CSV : {os.path.relpath(p1, ROOT)}")
    print(f"XLSX: {os.path.relpath(p2, ROOT) if p2 else 'openpyxl tidak tersedia'}")
